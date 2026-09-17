import openpyxl
from datetime import datetime
from openpyxl.styles import Font, PatternFill

FILE = 'Comprehensive_Salons_and_Staff_Register_EN_with_Payroll_12_REVIEWED.xlsx'
wb = openpyxl.load_workbook(FILE, data_only=False)
FONT_NAME = 'Arial'
normal_font = Font(name=FONT_NAME, size=10)
bold_font = Font(name=FONT_NAME, bold=True, size=10)

# ---------- 1. Add confirmed entries to MOHRE Verification (establishment 1278237) ----------
ws = wb['MOHRE Verification']
legend_row = None
for r in range(1, ws.max_row + 1):
    if ws.cell(row=r, column=1).value == 'Legend:':
        legend_row = r
        break

new_rows = [
    ('__HEADER__', 'Individual work-permit documents received 17/09/2026 for 2 more THE LAB - HQ people - '
                    'establishment 1278237 again. Mohammed Ahmed Abdallah Aljuneidi was previously categorized '
                    '"No Work Permit" (internal notes described him as a business partner) - this RENEWED '
                    'contract and work permit card show he does hold a genuine MOHRE work permit; correcting '
                    'his status to Confirmed.'),
    ('THE LAB - HQ', 'SALWAN BASHAR AL HAKIM', '✓ Confirmed', '122606279', '24/Dec/2026',
     'Work permit card and RENEWED contract on file (effective 25/12/2024) - establishment 1278237. '
     'Original hire date not on file - only the renewal contract was provided.'),
    ('THE LAB - HQ', 'MOHAMMED AHMED ABDALLAH ALJUNEIDI', '✓ Confirmed', '134843801', '09/Jan/2028',
     'Work permit card and RENEWED contract on file (effective 12/01/2026, salary AED 40,000 total matching the '
     'business-partner note already on file) - establishment 1278237. Status corrected from "No Work Permit" - '
     'he does hold a genuine MOHRE work permit. Original hire date not on file.'),
]

insert_at = legend_row
ws.insert_rows(insert_at, amount=len(new_rows) + 1)
r = insert_at
for row in new_rows:
    if row[0] == '__HEADER__':
        ws.cell(row=r, column=1, value=row[1]).font = bold_font
        r += 1
        continue
    salon, name, status, card, exp, detail = row
    ws.cell(row=r, column=1, value=salon)
    ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=3, value=status)
    ws.cell(row=r, column=4, value=card)
    ws.cell(row=r, column=5, value=exp)
    ws.cell(row=r, column=6, value=detail)
    for c in range(1, 7):
        ws.cell(row=r, column=c).font = normal_font
    r += 1
print(f'Inserted {len(new_rows)} rows into MOHRE Verification (before row {insert_at}).')

# ---------- 2. Notes for renewal-only people ----------
ws_hq = wb['THE LAB - HQ']
notes_append_hq = {
    'SALWAN BASHAR AL HAKIM': 'Renewed contract on file, effective 25/12/2024 - original hire date not on file '
                               '(only the renewal was provided); do not use this date for gratuity.',
    'MOHAMMED AHMED ABDALLAH ALJUNEIDI': 'Renewed contract on file, effective 12/01/2026 - work permit confirmed '
                                          '(status corrected from "No Work Permit"); original hire date not on file.',
}
for r in range(5, ws_hq.max_row + 1):
    name = ws_hq.cell(row=r, column=3).value
    if not name:
        continue
    name_u = str(name).upper().strip()
    if name_u in notes_append_hq:
        notes_cell = ws_hq.cell(row=r, column=30)
        existing = notes_cell.value or ''
        notes_cell.value = (existing + (' ' if existing else '') + notes_append_hq[name_u]).strip()
        notes_cell.font = normal_font

# ---------- 3. Gladys Balbiran Rampas - left the company (per salon manager, 17/09/2026) ----------
LEFT_FILL = PatternFill('solid', fgColor='D9D9D9')
for r in range(5, ws_hq.max_row + 1):
    name = ws_hq.cell(row=r, column=3).value
    if name and str(name).upper().strip() == 'GLADYS BALBIRAN RAMPAS':
        raw = str(name)
        if not raw.strip().startswith('🔴'):
            ws_hq.cell(row=r, column=3, value=f'🔴 LEFT — {raw}')
        action_cell = ws_hq.cell(row=r, column=7, value='Left Company')
        action_cell.font = normal_font
        notes_cell = ws_hq.cell(row=r, column=30)
        existing = notes_cell.value or ''
        note = ('Confirmed LEFT the company by the salon manager (17/09/2026). No MOHRE work permit card or '
                'labour contract was ever received for her - if she was working without one, this should be '
                'closed out with PRO/legal rather than pursued further; no gratuity/Join Date applicable.')
        notes_cell.value = (existing + (' ' if existing else '') + note).strip()
        notes_cell.font = normal_font
        for c in range(1, ws_hq.max_column + 1):
            ws_hq.cell(row=r, column=c).fill = LEFT_FILL
        print('Marked Gladys Balbiran Rampas as LEFT at row', r)

wb.save(FILE)
print('Saved.')
